#!/usr/bin/env python3
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pathlib import Path

root = Path(__file__).resolve().parent
book = root.parent / 'book' / 'chapters' / '04-crottys-house-1832.md'
if not book.exists():
    book = root / 'book' / 'chapters' / '04-crottys-house-1832.md'
raw = book.read_text()
if raw.startswith('---'):
    end = raw.find('\n---', 3)
    raw = raw[end + 4:].lstrip('\n')
paras = []
buf = []
skip_place = True
for line in raw.splitlines():
    if line.startswith('## Sources'):
        break
    if line.startswith('# '):
        continue
    if line.startswith('!['):
        continue
    if skip_place and line.startswith('*') and line.endswith('*') and line.count('*') == 2:
        skip_place = False
        continue
    if line.strip() == '---':
        continue
    if line.strip() == '':
        if buf:
            paras.append(' '.join(buf))
            buf = []
        continue
    buf.append(line.strip())
if buf:
    paras.append(' '.join(buf))

def chars_per_line(size_pt, width_in=8.45):
    return max(28, int(width_in * 72 / (size_pt * 0.46)))

def estimate_body_pt(paras, size, space_after):
    cpl = chars_per_line(size)
    lines = 0
    for para in paras:
        lines += max(1, (len(para) + cpl - 1) // cpl)
    return lines * (size * 1.16) + max(0, len(paras) - 1) * space_after

budget = 12.35 * 72
chosen = 8.0
gap = 3
for size in (14, 13, 12, 11.5, 11, 10.5, 10, 9.5, 9, 8.5, 8):
    gap = 8 if size >= 12 else 6 if size >= 10 else 4 if size >= 9 else 3
    if estimate_body_pt(paras, size, gap) <= budget:
        chosen = size
        break
else:
    chosen = 8
    gap = 2

prs = Presentation()
prs.slide_width = Inches(20)
prs.slide_height = Inches(14.25)
prs.core_properties.author = 'MJF'
prs.core_properties.title = "Crotty's House, 1832"
slide = prs.slides.add_slide(prs.slide_layouts[6])
bg = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(20), Inches(14.25))
bg.fill.solid()
bg.fill.fore_color.rgb = RGBColor(0xED, 0xE4, 0xD4)
bg.line.fill.background()
hair = slide.shapes.add_shape(1, Inches(9.15), Inches(0.35), Inches(0.015), Inches(13.55))
hair.fill.solid()
hair.fill.fore_color.rgb = RGBColor(0xC4, 0xB4, 0x9A)
hair.line.fill.background()
box = slide.shapes.add_textbox(Inches(0.45), Inches(0.32), Inches(8.5), Inches(13.6))
tf = box.text_frame
tf.word_wrap = True

def add_run(p, text, size, bold=False, italic=False, color=(0x2A, 0x22, 0x18)):
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = 'Palatino Linotype'
    run.font.color.rgb = RGBColor(*color)

p = tf.paragraphs[0]
add_run(p, 'MJF', 9, color=(0x6B, 0x3A, 0x2A))
p = tf.add_paragraph()
add_run(p, "CROTTY'S HOUSE, 1832", 16, bold=True)
p = tf.add_paragraph()
add_run(p, 'Kilrush, County Clare', 12, italic=True, color=(0x6B, 0x3A, 0x2A))
p = tf.add_paragraph()
add_run(p, '', 4)
for para in paras:
    p = tf.add_paragraph()
    add_run(p, para, chosen)
    p.space_after = Pt(gap)
names = [
 '01-conacre-strip.jpg','02-paving-the-house.jpg','03-onto-the-highway.jpg','04-timber-left-down.jpg',
 '05-kilrush-oath.jpg','06-ennis-counter.jpg','07-quinpool-gate.jpg','08-courthouse-lets-out.jpg']
left0, top0, gutter = 9.35, 0.4, 0.12
cell_w = (19.7 - left0 - gutter) / 2
cell_h = (13.95 - top0 - 3 * gutter) / 4
for i, name in enumerate(names):
    col, row = i % 2, i // 2
    x = left0 + col * (cell_w + gutter)
    y = top0 + row * (cell_h + gutter)
    pic = root / 'panels' / name
    if pic.exists():
        slide.shapes.add_picture(str(pic), Inches(x), Inches(y), width=Inches(cell_w))
out = root / 'Crottys_House_1832.pptx'
prs.save(str(out))
print('wrote', out, out.stat().st_size, 'body', chosen, 'gap', gap, 'paras', len(paras))
