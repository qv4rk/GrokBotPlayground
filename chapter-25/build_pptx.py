#!/usr/bin/env python3
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pathlib import Path
import re

root = Path(__file__).resolve().parent
book = root.parent / "book" / "chapters" / "25-two-letters-downing-street-1921.md"
if not book.exists():
    book = root / "book" / "chapters" / "25-two-letters-downing-street-1921.md"
raw = book.read_text()
if raw.startswith("---"):
    end = raw.find("\n---", 3)
    raw = raw[end + 4 :].lstrip("\n")
raw = re.sub(r"^#\s+[^\n]+\n+", "", raw)
raw = re.sub(r"^!\[[^\]]*\]\([^)]+\)\n+", "", raw)
raw = re.sub(r"^\*[^*]+\*\n+", "", raw)
idx = raw.find("\n## Sources")
if idx >= 0:
    raw = raw[:idx]
paras = []
for block in re.split(r"\n\s*\n", raw):
    t = " ".join(block.split())
    if t and t != "---" and not t.startswith("#"):
        paras.append(t)
prs = Presentation()
prs.slide_width = Inches(20)
prs.slide_height = Inches(14.25)
slide = prs.slides.add_slide(prs.slide_layouts[6])
bg = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(20), Inches(14.25))
bg.fill.solid()
bg.fill.fore_color.rgb = RGBColor(0xED, 0xE4, 0xD4)
bg.line.fill.background()
hair = slide.shapes.add_shape(1, Inches(9.15), Inches(0.35), Inches(0.015), Inches(13.55))
hair.fill.solid()
hair.fill.fore_color.rgb = RGBColor(0xC4, 0xB4, 0x9A)
hair.line.fill.background()
box = slide.shapes.add_textbox(Inches(0.45), Inches(0.35), Inches(8.5), Inches(13.55))
tf = box.text_frame
tf.word_wrap = True

def add_run(p, text, size, bold=False, italic=False, color=(0x2A, 0x22, 0x18)):
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = "Palatino Linotype"
    run.font.color.rgb = RGBColor(*color)

n = sum(len(p) for p in paras)
size = 8.5 if n < 5500 else 7.5
gap = 4 if n < 5500 else 3
p = tf.paragraphs[0]
add_run(p, "MJF", 9, color=(0x6B, 0x3A, 0x2A))
p = tf.add_paragraph()
add_run(p, "TWO LETTERS, DOWNING STREET, 1921", 16, bold=True)
p = tf.add_paragraph()
add_run(p, "London", 11, italic=True, color=(0x6B, 0x3A, 0x2A))
p = tf.add_paragraph()
add_run(p, "", 6)
for para in paras:
    p = tf.add_paragraph()
    add_run(p, para, size)
    p.space_after = Pt(gap)
names = [
    "01-two-letters.jpg",
    "02-euston-train.jpg",
    "03-holyhead-hull.jpg",
    "04-barton-holdout.jpg",
    "05-treaty-lamp.jpg",
    "06-collins-dawn.jpg",
    "07-earlsfort-house.jpg",
    "08-beal-na-blath.jpg",
]
left0, top0, gap_in = 9.35, 0.4, 0.12
cell_w = (19.7 - left0 - gap_in) / 2
cell_h = (13.95 - top0 - 3 * gap_in) / 4
for i, name in enumerate(names):
    col, row = i % 2, i // 2
    x = left0 + col * (cell_w + gap_in)
    y = top0 + row * (cell_h + gap_in)
    pic = root / "panels" / name
    if pic.exists():
        slide.shapes.add_picture(str(pic), Inches(x), Inches(y), width=Inches(cell_w))
out = root / "Two_Letters_Downing_Street_1921.pptx"
prs.save(str(out))
print("wrote", out, out.stat().st_size)
